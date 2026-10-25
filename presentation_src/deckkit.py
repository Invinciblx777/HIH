"""Drawing primitives for the AegisLab deck: native, editable PowerPoint shapes placed on the official Hack in Hills template.

Adapted from the helpers used for the earlier submission deck (same palette, fonts and layout grid, 20 x 11.25 in)."""
import copy
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt


INK, NAVY, ORANGE, ORANGE_SOFT = "1D1D1F", "13233A", "E8772E", "FCEBDD"
SLATE, SLATE_SOFT, LINE, MIST, WHITE = "5B6B7F", "EEF1F5", "C9D2DE", "F6F8FA", "FFFFFF"
RED, RED_SOFT, BLUE, BLUE_SOFT, GREEN, GREEN_SOFT = "C2453B", "FBE7E5", "2F6DB5", "E3EDF9", "2E8B62", "E2F3EA"
AMBER, AMBER_SOFT = "9A6200", "FFF1D6"
HEAD, BODY = "Montserrat", "IBM Plex Sans"


def rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------ primitives
def _font(run, size, bold=False, color=INK, font=BODY, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def text(slide, x, y, w, h, paras, size=20, color=INK, bold=False, font=BODY, align="l", anchor="t",
         fill=None, line=None, margin=0.0, spacing=None, name=None, shape=None, radius=None, line_w=1.25):
    """paras: str | list of paragraphs; a paragraph is str or list of (text, opts) runs."""
    if shape is None:
        shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius
    if name:
        shp.name = name
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    elif shape is not None:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    elif shape is not None:
        shp.line.fill.background()
    if shape is not None:
        shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            para.space_after = Pt(spacing)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("font", font),
                  o.get("italic", False))
    return shp


def box(slide, x, y, w, h, title, subs=(), fill=WHITE, line=SLATE, tcol=INK, tsize=21, ssize=16, line_w=2,
        dash=False, name=None, align="c"):
    paras = [[(title, {"bold": True, "font": HEAD, "size": tsize, "color": tcol})]]
    paras += [[(s, {"size": ssize, "color": SLATE})] for s in subs]
    shp = text(slide, x, y, w, h, paras, align=align, anchor="m", fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
               radius=0.12, margin=0.08, line_w=line_w, name=name)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shp


def frame(slide, x, y, w, h, color, dash=True, fill=None, line_w=2, name=None, radius=0.04):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(line_w)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def arrow(slide, pts, color, width=2.5, dash=False, head=True, both=False, name=None):
    """Polyline connector (freeform) with arrowheads; pts in inches."""
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    shp.shadow.inherit = False
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = shp.line._get_or_add_ln()
    if both:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    if head:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    if name:
        shp.name = name
    return shp


def label(slide, x, y, s, color, size=15, w=None, align="c", bg=WHITE):
    w = w or len(s) * size * 0.0082 + 0.2
    return text(slide, x - (w / 2 if align == "c" else 0), y, w, size / 72 * 1.35, [[(s, {"bold": True})]], size=size,
                color=color, align=align, anchor="m", fill=bg, shape=MSO_SHAPE.RECTANGLE if bg else None)


def chip(slide, x, y, s, fg, bg, size=13):
    w = len(s) * size * 0.0098 + 0.32
    text(slide, x, y, w, size / 72 + 0.16, [[(s.upper(), {"bold": True})]], size=size, color=fg, fill=bg, align="c",
         anchor="m", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, margin=0.02)
    return w


def circle(slide, x, y, d, s, fill, size=16):
    return text(slide, x, y, d, d, [[(s, {"bold": True, "font": HEAD})]], size=size, color=WHITE, fill=fill,
                align="c", anchor="m", shape=MSO_SHAPE.OVAL, margin=0)


def header(slide, kicker, headline):
    text(slide, 0.7, 1.55, 18.6, 0.42, kicker, size=20, bold=True, color=ORANGE, font=HEAD, name="Section")
    text(slide, 0.7, 1.98, 18.6, 0.8, headline, size=36, bold=True, color=INK, font=HEAD, name="Title")


def picture(slide, path, x, y, w=None, h=None, border=LINE):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic.line.color.rgb = rgb(border)
    pic.line.width = Pt(1.25)
    return pic


def clear(slide):
    for shp in list(slide.shapes):
        if shp.shape_type != MSO_SHAPE_TYPE.PICTURE:
            shp._element.getparent().remove(shp._element)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s.strip()

